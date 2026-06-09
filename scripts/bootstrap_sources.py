from __future__ import annotations

import csv
import re
import subprocess
from datetime import date, datetime
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
REF = ROOT / "references" / "originals"
SOURCE = ROOT / "source"
DATA = ROOT / "data" / "mentors-2025"


PDF_SOURCES = [
    {
        "pdf": "welcome-to-ucas-original.pdf",
        "target": SOURCE / "chapters" / "01-first-lesson.tex",
        "title": "来到国科大的第一课",
        "attribution": "本章逆向整理自原 PDF《来到国科大的第一课！》，原作者：试管，中国科学院大学本科部 23 级学生。",
    },
    {
        "pdf": "course-selection-original.pdf",
        "target": SOURCE / "chapters" / "02-course-selection.tex",
        "title": "如何选课",
        "attribution": "本章逆向整理自原 PDF《如何选课？》，原作者：试管，中国科学院大学本科部 23 级学生。21、22 级学长们和 23 级同学曾提出宝贵意见。",
    },
    {
        "pdf": "enterprise-wechat-guide.pdf",
        "target": SOURCE / "chapters" / "03-enterprise-wechat.tex",
        "title": "企业号关注流程及常见问题",
        "attribution": "本章逆向整理自原 PDF《中国科学院大学企业号关注流程及常见问题解决方法》。原 PDF 未标注个人作者；截图和二维码请以归档原件为准。",
    },
]


MENTOR_SLUGS = [
    "mathematics",
    "physics",
    "chemistry",
    "biology",
    "materials-science",
    "computer-science",
    "astronomy",
    "electronic-information",
    "environmental-science",
    "mechanics",
    "human-geography",
    "electrical-engineering",
    "cyberspace-security",
    "artificial-intelligence",
    "psychology",
]


LATEX_SPECIALS = {
    "\\": r"\textbackslash{}",
    "&": r"\&",
    "%": r"\%",
    "$": r"\$",
    "#": r"\#",
    "_": r"\_",
    "{": r"\{",
    "}": r"\}",
    "~": r"\textasciitilde{}",
    "^": r"\textasciicircum{}",
}

CIRCLED_DIGITS = str.maketrans(
    {
        "①": "(1)",
        "②": "(2)",
        "③": "(3)",
        "④": "(4)",
        "⑤": "(5)",
        "⑥": "(6)",
        "⑦": "(7)",
        "⑧": "(8)",
        "⑨": "(9)",
    }
)


def escape_tex(text: str) -> str:
    text = text.translate(CIRCLED_DIGITS)
    return "".join(LATEX_SPECIALS.get(ch, ch) for ch in text)


def run_pdftotext(pdf_path: Path) -> str:
    result = subprocess.run(
        ["pdftotext", "-raw", "-enc", "UTF-8", str(pdf_path), "-"],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    return result.stdout.decode("utf-8", errors="replace")


def should_skip_line(line: str) -> bool:
    stripped = line.strip()
    return (
        stripped.startswith("作者：试管，中国科学院大学本科部")
        or stripped.startswith("21、22 级学长们和 23 级同学提出了宝贵的意见")
    )


def is_new_paragraph_marker(line: str) -> bool:
    return bool(
        re.match(r"^([一二三四五六七八九十]+、|\d+[、.]|[①②③④⑤⑥⑦⑧⑨])", line)
        or re.match(r"^(方法|可能性)[一二三四五六七八九十]", line)
        or line.startswith(("确认方式：", "处理办法：", "备注：", "注："))
    )


def is_complete_enough(text: str) -> bool:
    if len(text) < 45:
        return False
    return text.endswith(("。", "！", "？", "；", "）", ")", "：", ":"))


def is_standalone_heading_line(line: str) -> bool:
    return bool(
        len(line) <= 38
        and (
            re.match(r"^[一二三四五六七八九十]+、", line)
            or re.match(r"^\d+[、.]", line)
            or re.match(r"^(方法|可能性)[一二三四五六七八九十]：?$", line)
        )
    )


def normalize_pdf_lines(text: str) -> list[str]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").replace("\f", "\n")
    lines = []
    for raw_line in normalized.splitlines():
        line = re.sub(r"\s+", " ", raw_line.strip())
        if not line or should_skip_line(line):
            continue
        lines.append(line)

    combined: list[str] = []
    index = 0
    while index < len(lines):
        line = lines[index]
        if re.fullmatch(r"\d+[、.]", line) and index + 1 < len(lines):
            combined.append(line + lines[index + 1])
            index += 2
            continue
        combined.append(line)
        index += 1
    return combined


def paragraphs_from_pdf_text(text: str) -> list[str]:
    paragraphs: list[str] = []
    current = ""

    def flush() -> None:
        nonlocal current
        if current.strip():
            paragraphs.append(current.strip())
        current = ""

    for line in normalize_pdf_lines(text):
        if is_new_paragraph_marker(line):
            flush()
            current = line
            if is_standalone_heading_line(line):
                flush()
                continue
        elif current:
            current += line
        else:
            current = line

        if is_complete_enough(current):
            flush()

    flush()
    return [p for p in paragraphs if p]


def split_heading(paragraph: str) -> tuple[str | None, str]:
    chinese = re.match(r"^([一二三四五六七八九十]+)、(.+?)(。|：|:)(.*)$", paragraph)
    if chinese:
        return chinese.group(2).strip(), chinese.group(4).strip()

    chinese_without_punctuation = re.match(r"^([一二三四五六七八九十]+)、(.+)$", paragraph)
    if chinese_without_punctuation and len(chinese_without_punctuation.group(2)) <= 28:
        return chinese_without_punctuation.group(2).strip("：:"), ""

    numbered = re.match(r"^(\d+)[、.](.+?)(。|：|:)(.*)$", paragraph)
    if numbered:
        return numbered.group(2).strip(), numbered.group(4).strip()

    numbered_without_punctuation = re.match(r"^(\d+)[、.](.+)$", paragraph)
    if numbered_without_punctuation and len(numbered_without_punctuation.group(2)) <= 36:
        return numbered_without_punctuation.group(2).strip("：:"), ""

    return None, paragraph


def write_pdf_chapter(source: dict[str, object]) -> None:
    pdf_path = REF / str(source["pdf"])
    text = run_pdftotext(pdf_path)
    paragraphs = paragraphs_from_pdf_text(text)

    lines = [
        f"% Auto-generated from references/originals/{source['pdf']} by scripts/bootstrap_sources.py.",
        "% The extraction is intentionally conservative; revise wording and structure during editing.",
        f"\\chapter{{{escape_tex(str(source['title']))}}}",
        f"\\sourceattribution{{{escape_tex(str(source['attribution']))}}}",
        "",
    ]

    first_paragraph = True
    for paragraph in paragraphs:
        if first_paragraph and paragraph.rstrip("！?？") in {
            "来到国科大的第一课",
            "来到国科大的第一课！",
            "如何选课",
            "如何选课？",
            "中国科学院大学企业号关注流程及常见问题解决方法",
        }:
            first_paragraph = False
            continue
        first_paragraph = False

        heading, remainder = split_heading(paragraph)
        if heading:
            lines.append(f"\\section{{{escape_tex(heading)}}}")
            lines.append("")
            if remainder:
                lines.append(escape_tex(remainder))
                lines.append("")
            continue

        lines.append(escape_tex(paragraph))
        lines.append("")

    if source["pdf"] == "enterprise-wechat-guide.pdf":
        lines.extend(
            [
                "\\begin{editingnote}",
                "原 PDF 包含二维码和多张微信、企业号操作截图。当前章节先保留文字逆向结果；排版时请对照 \\path{references/originals/enterprise-wechat-guide.pdf} 补图，并核对 2026 级实际企业号认证流程。",
                "\\end{editingnote}",
                "",
            ]
        )

    source["target"].write_text("\n".join(lines), encoding="utf-8")


def write_military_training_appendix() -> None:
    doc_path = REF / "military-training-schedule-2023.docx"
    document = Document(str(doc_path))

    title = "军训作息参考"
    note = ""
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text and title not in text:
            note = text
            break

    rows: list[list[str]] = []
    if document.tables:
        for row in document.tables[0].rows:
            rows.append([cell.text.strip().replace("\n", " / ") for cell in row.cells])

    lines = [
        "% Auto-generated from references/originals/military-training-schedule-2023.docx by scripts/bootstrap_sources.py.",
        "\\chapter{军训作息参考}",
        "\\sourceattribution{本附录整理自归档资料《23级军训作息安排.docx》。该资料为 2023 级参考作息，2026 级安排请以后续官方通知为准。}",
        "",
    ]
    if note:
        lines.append(escape_tex(note))
        lines.append("")

    lines.extend(
        [
            "\\begin{longtable}{>{\\raggedright\\arraybackslash}p{0.28\\linewidth}>{\\raggedright\\arraybackslash}p{0.62\\linewidth}}",
            "\\toprule",
        ]
    )
    for index, row in enumerate(rows):
        if len(row) < 2:
            continue
        left, right = escape_tex(row[0]), escape_tex(row[1])
        if index == 0:
            lines.append(f"\\textbf{{{left}}} & \\textbf{{{right}}} \\\\")
            lines.append("\\midrule")
        else:
            lines.append(f"{left} & {right} \\\\")
    lines.extend(["\\bottomrule", "\\end{longtable}", ""])

    (SOURCE / "appendices" / "military-training-schedule.tex").write_text(
        "\n".join(lines), encoding="utf-8"
    )


def clean_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def export_mentor_list() -> list[tuple[str, str, int]]:
    DATA.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(
        str(REF / "academic-mentor-list-2025.xlsx"), read_only=True, data_only=True
    )
    exported: list[tuple[str, str, int]] = []

    for index, worksheet in enumerate(workbook.worksheets, start=1):
        slug = MENTOR_SLUGS[index - 1] if index - 1 < len(MENTOR_SLUGS) else f"sheet-{index:02d}"
        filename = f"{index:02d}-{slug}.csv"
        path = DATA / filename
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle, lineterminator="\n")
            for row in worksheet.iter_rows(values_only=True):
                writer.writerow([clean_cell(value) for value in row])
        exported.append((worksheet.title.strip(), filename, max(worksheet.max_row - 1, 0)))

    return exported


def write_mentor_docs(exported: list[tuple[str, str, int]]) -> None:
    readme_lines = [
        "# 2025 年度本科生学业导师名单",
        "",
        "本目录由 `references/originals/academic-mentor-list-2025.xlsx` 导出，便于多人协作时查看差异和检索。",
        "导出脚本：`scripts/bootstrap_sources.py`。",
        "",
        "| 专业方向 | CSV | 记录数 |",
        "| --- | --- | ---: |",
    ]
    appendix_lines = [
        "\\chapter{学业导师名单参考}",
        "\\sourceattribution{本附录索引整理自归档资料《2025年度本科生学业导师名单》。2026 级实际导师名单请以本科部发布版本为准。}",
        "",
        "完整表格已按专业方向导出到 \\path{data/mentors-2025/}，方便在 Git 中审阅差异。本手册正文只保留索引，不直接排入完整名单。",
        "",
        "\\begin{longtable}{>{\\raggedright\\arraybackslash}p{0.42\\linewidth}>{\\raggedright\\arraybackslash}p{0.38\\linewidth}r}",
        "\\toprule",
        "\\textbf{专业方向} & \\textbf{CSV 文件} & \\textbf{记录数} \\\\",
        "\\midrule",
    ]

    for title, filename, count in exported:
        readme_lines.append(f"| {title} | `{filename}` | {count} |")
        appendix_lines.append(
            f"{escape_tex(title)} & \\path{{data/mentors-2025/{filename}}} & {count} \\\\"
        )

    appendix_lines.extend(["\\bottomrule", "\\end{longtable}", ""])

    (DATA / "README.md").write_text("\n".join(readme_lines) + "\n", encoding="utf-8")
    (SOURCE / "appendices" / "mentor-list.tex").write_text(
        "\n".join(appendix_lines), encoding="utf-8"
    )


def main() -> None:
    (SOURCE / "chapters").mkdir(parents=True, exist_ok=True)
    (SOURCE / "appendices").mkdir(parents=True, exist_ok=True)
    DATA.mkdir(parents=True, exist_ok=True)

    for pdf_source in PDF_SOURCES:
        write_pdf_chapter(pdf_source)
    write_military_training_appendix()
    exported = export_mentor_list()
    write_mentor_docs(exported)


if __name__ == "__main__":
    main()
